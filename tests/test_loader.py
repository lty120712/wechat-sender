# -*- coding: utf-8 -*-
import os
import sys
import tempfile
import unittest

sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))

from src.loader import load_messages, load_timed_messages, MessagePicker, MessageItem, TimedMessageItem


class TestLoadMessagesFromTxt(unittest.TestCase):

    def _write_tmp(self, content: str) -> str:
        f = tempfile.NamedTemporaryFile(mode="w", suffix=".txt", delete=False, encoding="utf-8")
        f.write(content)
        f.close()
        return f.name

    def test_single_line(self):
        fname = self._write_tmp("msg1\nmsg2\nmsg3\n")
        try:
            msgs = load_messages(fname, "")
            self.assertEqual(len(msgs), 3)
            self.assertEqual(msgs[0].title, "msg1")
            self.assertEqual(msgs[0].body, "msg1")
        finally:
            os.unlink(fname)

    def test_block_format(self):
        fname = self._write_tmp("---\nline1\nline2\n---\nblock2\n---\nblock3\nline4\n")
        try:
            msgs = load_messages(fname, "")
            self.assertEqual(len(msgs), 3)
            self.assertEqual(msgs[0].body, "line1\nline2")
            self.assertEqual(msgs[1].body, "block2")
            self.assertEqual(msgs[2].body, "block3\nline4")
        finally:
            os.unlink(fname)

    def test_comment_lines(self):
        fname = self._write_tmp("#comment\nmsg1\n;also comment\nmsg2\n")
        try:
            msgs = load_messages(fname, "")
            self.assertEqual(len(msgs), 2)
            self.assertEqual(msgs[0].body, "msg1")
            self.assertEqual(msgs[1].body, "msg2")
        finally:
            os.unlink(fname)


class TestLoadMessagesFromExcel(unittest.TestCase):

    def _make_xlsx(self) -> str:
        import openpyxl
        f = tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False)
        f.close()
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "循环消息"
        ws.append(["消息标题", "消息体", "发送条件"])
        ws.append(["早安", "早上好", "1"])
        ws.append(["晚安", "晚安，好梦", "*"])
        wb.save(f.name)
        return f.name

    def test_load_excel(self):
        fname = self._make_xlsx()
        try:
            msgs = load_messages(fname, "")
            self.assertEqual(len(msgs), 2)
            self.assertEqual(msgs[0].title, "早安")
            self.assertEqual(msgs[0].body, "早上好")
            self.assertEqual(msgs[0].day_condition, "1")
            self.assertEqual(msgs[1].title, "晚安")
            self.assertEqual(msgs[1].body, "晚安，好梦")
            self.assertEqual(msgs[1].day_condition, "*")
        finally:
            os.unlink(fname)

    def test_empty_body_skipped(self):
        import openpyxl
        f = tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False)
        f.close()
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "循环消息"
        ws.append(["消息标题", "消息体", "发送条件"])
        ws.append(["only title", "", ""])
        ws.append(["valid", "real msg", ""])
        wb.save(f.name)
        try:
            msgs = load_messages(f.name, "")
            self.assertEqual(len(msgs), 1)
            self.assertEqual(msgs[0].body, "real msg")
        finally:
            os.unlink(f.name)

    def test_day_condition_loaded(self):
        import openpyxl
        f = tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False)
        f.close()
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "循环消息"
        ws.append(["消息标题", "消息体", "发送条件"])
        ws.append(["only_weekend", "周末消息", "6,7"])
        wb.save(f.name)
        try:
            msgs = load_messages(f.name, "")
            self.assertEqual(len(msgs), 1)
            self.assertEqual(msgs[0].day_condition, "6,7")
        finally:
            os.unlink(f.name)

    def test_invalid_condition_raises_with_row(self):
        import openpyxl
        f = tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False)
        f.close()
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "循环消息"
        ws.append(["消息标题", "消息体", "发送条件"])
        ws.append(["bad", "something", "abc"])
        wb.save(f.name)
        try:
            with self.assertRaises(ValueError) as ctx:
                load_messages(f.name, "")
            self.assertIn("第 2 行", str(ctx.exception))
        finally:
            os.unlink(f.name)

    def test_empty_enabled_is_disabled(self):
        import openpyxl
        f = tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False)
        f.close()
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "循环消息"
        ws.append(["消息标题", "消息体", "发送条件", "是否启用"])
        ws.append(["msg", "hello", "*", ""])
        wb.save(f.name)
        try:
            msgs = load_messages(f.name, "")
            self.assertEqual(len(msgs), 1)
            self.assertFalse(msgs[0].enabled)
        finally:
            os.unlink(f.name)


class TestLoadTimedMessages(unittest.TestCase):

    def _make_xlsx(self) -> str:
        import openpyxl
        f = tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False)
        f.close()
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "循环消息"
        ws.append(["消息标题", "消息体"])
        ws.append(["normal", "hello"])
        ws2 = wb.create_sheet("到点消息")
        ws2.append(["时间", "消息标题", "消息体", "发送条件"])
        ws2.append(["14:00", "下午", "下午好", "1"])
        ws2.append(["18:00", "", "下班了", ""])
        wb.save(f.name)
        return f.name

    def test_load_timed(self):
        fname = self._make_xlsx()
        try:
            items = load_timed_messages(fname, "")
            self.assertEqual(len(items), 2)
            self.assertEqual(items[0].time, "14:00")
            self.assertEqual(items[0].title, "下午")
            self.assertEqual(items[0].body, "下午好")
            self.assertEqual(items[0].day_condition, "1")
            self.assertEqual(items[1].time, "18:00")
            self.assertEqual(items[1].body, "下班了")
            self.assertEqual(items[1].day_condition, "")
        finally:
            os.unlink(fname)

    def test_no_timed_sheet_returns_empty(self):
        import openpyxl
        f = tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False)
        f.close()
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "循环消息"
        ws.append(["消息标题", "消息体"])
        wb.save(f.name)
        try:
            items = load_timed_messages(f.name, "")
            self.assertEqual(items, [])
        finally:
            os.unlink(f.name)

    def test_txt_returns_empty(self):
        with tempfile.NamedTemporaryFile(suffix=".txt", mode="w", delete=False, encoding="utf-8") as f:
            f.write("hello\n")
            fname = f.name
        try:
            items = load_timed_messages(fname, "")
            self.assertEqual(items, [])
        finally:
            os.unlink(fname)

    def test_invalid_timed_raises_with_row(self):
        import openpyxl
        f = tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False)
        f.close()
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "循环消息"
        ws.append(["消息标题", "消息体"])
        ws.append(["x", "y"])
        ws2 = wb.create_sheet("到点消息")
        ws2.append(["时间", "消息标题", "消息体", "发送条件"])
        ws2.append(["14:00", "bad", "something", "8"])
        wb.save(f.name)
        try:
            with self.assertRaises(ValueError) as ctx:
                load_timed_messages(f.name, "")
            self.assertIn("第 2 行", str(ctx.exception))
        finally:
            os.unlink(f.name)


class TestMessagePicker(unittest.TestCase):

    def test_random(self):
        msgs = [MessageItem("a", "a"), MessageItem("b", "b"), MessageItem("c", "c")]
        p = MessagePicker(msgs, "random")
        picked = {p.pick().body for _ in range(100)}
        self.assertTrue(picked.issubset({"a", "b", "c"}))

    def test_sequential(self):
        msgs = [MessageItem("a", "a"), MessageItem("b", "b"), MessageItem("c", "c")]
        p = MessagePicker(msgs, "sequential")
        self.assertEqual(p.pick().body, "a")
        self.assertEqual(p.pick().body, "b")
        self.assertEqual(p.pick().body, "c")
        self.assertEqual(p.pick().body, "a")

    def test_filter_skips_non_matching(self):
        from datetime import datetime
        other = str((datetime.now().weekday() + 1) % 7 + 1)
        msgs = [MessageItem("x", "skip", day_condition=other)]
        p = MessagePicker(msgs, "sequential")
        self.assertIsNone(p.pick())

    def test_filter_picks_matching(self):
        from datetime import datetime
        msgs = [MessageItem("x", "match", day_condition=str(datetime.now().weekday() + 1))]
        p = MessagePicker(msgs, "sequential")
        m = p.pick()
        self.assertIsNotNone(m)
        self.assertEqual(m.body, "match")

    def test_filter_comma_list_matches(self):
        from datetime import datetime
        t = datetime.now().weekday() + 1
        msgs = [MessageItem("x", "comma", day_condition=f"{t},5,7")]
        p = MessagePicker(msgs, "sequential")
        m = p.pick()
        self.assertIsNotNone(m)
        self.assertEqual(m.body, "comma")

    def test_disabled_message_skipped(self):
        msgs = [MessageItem("x", "disabled", enabled=False)]
        p = MessagePicker(msgs, "sequential")
        self.assertIsNone(p.pick())

    def test_disabled_message_mixed(self):
        msgs = [
            MessageItem("a", "a", enabled=False),
            MessageItem("b", "b", enabled=True),
        ]
        p = MessagePicker(msgs, "sequential")
        m = p.pick()
        self.assertIsNotNone(m)
        self.assertEqual(m.body, "b")


if __name__ == "__main__":
    unittest.main()
