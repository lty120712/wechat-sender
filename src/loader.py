# -*- coding: utf-8 -*-
import os
import random
from dataclasses import dataclass
from typing import List, Optional

from src.utils import match_day_condition, validate_day_condition


@dataclass
class MessageItem:
    title: str
    body: str
    day_condition: str = ""
    enabled: bool = True


@dataclass
class TimedMessageItem:
    time: str
    title: str
    body: str
    day_condition: str = ""
    enabled: bool = True


@dataclass
class HolidayMessageItem:
    name: str
    month_day: str
    time: str
    body: str
    enabled: bool = True


def _decode(text: str) -> str:
    return text.replace("\r\n", "\n")


def _load_from_txt(filepath: str) -> List[MessageItem]:
    with open(filepath, "r", encoding="utf-8-sig") as f:
        lines = f.readlines()

    messages: List[MessageItem] = []
    has_block = any(line.strip() == "---" for line in lines)

    if has_block:
        block: List[str] = []
        for line in lines:
            stripped = line.strip()
            if stripped == "---":
                msg = "\n".join(block).strip()
                if msg:
                    title = block[0].strip() if block else ""
                    messages.append(MessageItem(title=title, body=_decode(msg)))
                block = []
                continue
            if stripped.startswith("#") or stripped.startswith(";"):
                continue
            block.append(line.rstrip("\r\n"))
        msg = "\n".join(block).strip()
        if msg:
            title = block[0].strip() if block else ""
            messages.append(MessageItem(title=title, body=_decode(msg)))
    else:
        for line in lines:
            msg = line.strip()
            if not msg or msg.startswith("#") or msg.startswith(";"):
                continue
            messages.append(MessageItem(title=msg, body=_decode(msg)))

    if not messages:
        raise ValueError("循环消息为空")
    return messages


def _load_from_excel(filepath: str) -> List[MessageItem]:
    import openpyxl

    wb = openpyxl.load_workbook(filepath)
    ws = wb["循环消息"]
    messages: List[MessageItem] = []
    for i, row in enumerate(ws.iter_rows(min_row=2, values_only=True), 2):
        if not row or len(row) < 2:
            continue
        title = str(row[0]).strip() if row[0] else ""
        body = str(row[1]).strip() if row[1] else ""
        day = str(row[2]).strip() if len(row) > 2 and row[2] else ""
        enabled = str(row[3]).strip().upper() == "Y" if len(row) > 3 else True
        if not body:
            continue
        try:
            day = validate_day_condition(day)
        except ValueError as exc:
            raise ValueError(f"[循环消息] 第 {i} 行：{exc}")
        messages.append(MessageItem(title=title, body=body, day_condition=day, enabled=enabled))
    if not messages:
        raise ValueError("循环消息为空：Excel [循环消息] 工作表中没有数据")
    return messages


def load_messages(source: str, base_dir: str) -> List[MessageItem]:
    source_path = source if os.path.isabs(source) else os.path.join(base_dir, source)
    if not os.path.exists(source_path):
        raise FileNotFoundError(f"消息文件不存在：{source_path}")

    ext = os.path.splitext(source_path)[1].lower()
    if ext == ".xlsx":
        return _load_from_excel(source_path)
    return _load_from_txt(source_path)


def load_timed_messages(source: str, base_dir: str) -> List[TimedMessageItem]:
    source_path = source if os.path.isabs(source) else os.path.join(base_dir, source)
    if not os.path.exists(source_path):
        return []

    ext = os.path.splitext(source_path)[1].lower()
    if ext != ".xlsx":
        return []

    import openpyxl

    wb = openpyxl.load_workbook(source_path)
    if "到点消息" not in wb.sheetnames:
        return []

    ws = wb["到点消息"]
    result: List[TimedMessageItem] = []
    for i, row in enumerate(ws.iter_rows(min_row=2, values_only=True), 2):
        if not row or len(row) < 3:
            continue
        t = str(row[0]).strip() if row[0] else ""
        title = str(row[1]).strip() if row[1] else ""
        body = str(row[2]).strip() if row[2] else ""
        day = str(row[3]).strip() if len(row) > 3 and row[3] else ""
        enabled = str(row[4]).strip().upper() == "Y" if len(row) > 4 else True
        if not t or not body:
            continue
        try:
            day = validate_day_condition(day)
        except ValueError as exc:
            raise ValueError(f"[到点消息] 第 {i} 行：{exc}")
        result.append(TimedMessageItem(time=t, title=title, body=body, day_condition=day, enabled=enabled))
    return result


def load_holiday_messages(source: str, base_dir: str) -> List[HolidayMessageItem]:
    source_path = source if os.path.isabs(source) else os.path.join(base_dir, source)
    if not os.path.exists(source_path):
        return []

    ext = os.path.splitext(source_path)[1].lower()
    if ext != ".xlsx":
        return []

    import openpyxl

    wb = openpyxl.load_workbook(source_path)
    if "节日消息" not in wb.sheetnames:
        return []

    ws = wb["节日消息"]
    result: List[HolidayMessageItem] = []
    for i, row in enumerate(ws.iter_rows(min_row=2, values_only=True), 2):
        if not row or len(row) < 4:
            continue
        name = str(row[0]).strip() if row[0] else ""
        month_day = str(row[1]).strip() if row[1] else ""
        t = str(row[2]).strip() if row[2] else ""
        body = str(row[3]).strip() if row[3] else ""
        enabled = str(row[4]).strip().upper() == "Y" if len(row) > 4 else True
        if not name or not month_day or not t or not body:
            continue
        import re
        if not re.match(r"^\d{2}-\d{2}$", month_day):
            raise ValueError(f"[节日消息] 第 {i} 行：月日 '{month_day}' 格式错误，应如 01-01")
        result.append(HolidayMessageItem(name=name, month_day=month_day, time=t, body=body, enabled=enabled))
    return result


class MessagePicker:
    def __init__(self, messages: List[MessageItem], mode: str = "sequential"):
        if not messages:
            raise ValueError("消息列表不能为空")
        self._messages = messages
        self._mode = mode.lower()
        self._index = 0

    @property
    def count(self) -> int:
        return len(self._messages)

    def _eligible(self, m: MessageItem) -> bool:
        return m.enabled and match_day_condition(m.day_condition)

    def pick(self) -> Optional[MessageItem]:
        if self._mode == "random":
            eligible = [m for m in self._messages if self._eligible(m)]
            if not eligible:
                return None
            return random.choice(eligible)
        for _ in range(len(self._messages)):
            msg = self._messages[self._index % len(self._messages)]
            self._index += 1
            if self._eligible(msg):
                return msg
        return None

    def reset(self) -> None:
        self._index = 0
